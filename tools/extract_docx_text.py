from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def extract_docx_text(path: Path) -> str:
    doc = Document(str(path))
    lines: list[str] = []

    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            lines.append(t)

    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip().replace("\n", " ").strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    return "\n".join(lines).strip() + "\n"


def main() -> None:
    # Ensure Windows console/file redirection uses UTF-8 for SRS symbols.
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # py3.7+
    except Exception:
        pass
    path = Path(r"D:\mission planner\SRS_Drone_GCS_v3 (1).docx")
    print(extract_docx_text(path), end="")


if __name__ == "__main__":
    main()

